import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.models.schemas import ExtractionOut, FieldValueOut, Job, PageContent, PageTriage, RecordGroupOut, TriageResult
from app.services.export_service import ExportService
from app.services.field_catalog import get_catalog
from app.services.normalization_service import NormalizationService


def _job():
    tri = TriageResult(document_type="BENEFIT_SCHEDULE", reasoning="t",
                       pages=[PageTriage(page=1, role="BENEFIT_TABLE", is_relevant=True, plan_columns=["แผน 001", "แผน 002"])])
    out = ExtractionOut(case_fields=[], groups=[
        RecordGroupOut(group_key="001", fields=[FieldValueOut(field_code="BEN_IPD_RB", value="1500", evidence="ค่าห้อง", page=1)]),
        RecordGroupOut(group_key="002", fields=[FieldValueOut(field_code="BEN_IPD_RB", value="3000", evidence="ค่าห้อง", page=1)]),
    ])
    result = NormalizationService(get_catalog()).normalize("BENEFIT_SCHEDULE", out, [PageContent(page_no=1, source_file="a.jpg", markdown="ค่าห้อง 1,500 3,000")], tri, [], {"ocr": 1}, 2)
    job = Job(job_id="abc", case_name="Case 1", filenames=["a.jpg"], status="completed", result=result)
    return job


def test_export_matrix_workbook():
    data = ExportService().build_workbook(_job())
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Benefit schedule", "Field inspector", "Run info"]
    ws = wb["Benefit schedule"]
    headers = [c.value for c in ws[5]]
    assert headers[:3] == ["Benefit", "As printed (Thai)", "Field code"] and "แผน 001" in headers[3]
    codes = [row[2].value for row in ws.iter_rows(min_row=6) if row[2].value]
    assert "BEN_IPD_RB" in codes
    assert "BEN_GTL" not in codes


def test_health_catalog_and_patch_flow():
    client = TestClient(app)
    assert client.get("/health").json()["catalog"]["field_count"] > 100
    cat = client.get("/api/v1/catalog").json()
    assert set(cat["doc_types"]) == {"BENEFIT_SCHEDULE", "CLAIMS", "CENSUS"}

    from app.services.pipeline_service import get_pipeline_service
    job = _job()
    get_pipeline_service().jobs[job.job_id] = job
    res = client.patch(f"/api/v1/documents/{job.job_id}/fields", json={"overrides": [{"group_key": "001", "field_code": "BEN_DEN", "value": "4,000"}]})
    assert res.status_code == 200
    g = next(g for g in res.json()["result"]["groups"] if g["group_key"] == "001")
    f = next(f for f in g["fields"] if f["field_code"] == "BEN_DEN")
    assert f["value"] == 4000 and f["edited"] is True and f["status"] == "extracted"
    exp = client.get(f"/api/v1/documents/{job.job_id}/export")
    assert exp.status_code == 200 and exp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert client.get("/api/v1/documents/nope").status_code == 404


